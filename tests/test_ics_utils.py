# =============================================================
#  Transkription_Notes_Pipeline - tests/test_ics_utils.py
#  Unit tests for ics_utils.py: the .ics VEVENT parser (single-event,
#  task #93 extension, and whole-calendar, task #94), the labeled-fields
#  .txt/.docx parser, and the parse_meeting_info_file dispatcher.
# =============================================================

import sys
from datetime import datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import ics_utils


# -----------------------------------------------------------------
# parse_ics
# -----------------------------------------------------------------

ICS_TEMPLATE = """BEGIN:VCALENDAR
VERSION:2.0
BEGIN:VEVENT
SUMMARY:{summary}
DTSTART:{dtstart}
{description_line}END:VEVENT
END:VCALENDAR
"""


class TestParseIcs:
    def _write_ics(self, tmp_path, summary="Weekly Sync", dtstart="20260704T090000",
                    description=None):
        desc_line = f"DESCRIPTION:{description}\n" if description else ""
        content = ICS_TEMPLATE.format(summary=summary, dtstart=dtstart, description_line=desc_line)
        p = tmp_path / "invite.ics"
        p.write_text(content, encoding="utf-8")
        return p

    def test_title_and_date(self, tmp_path):
        p = self._write_ics(tmp_path)
        result = ics_utils.parse_ics(str(p))
        assert result["title"] == "Weekly Sync"
        assert result["date"] == "2026-07-04"

    def test_description_becomes_comments(self, tmp_path):
        p = self._write_ics(tmp_path, description="Bring the Q3 numbers.")
        result = ics_utils.parse_ics(str(p))
        assert result["comments"] == "Bring the Q3 numbers."

    def test_no_description_yields_blank_comments(self, tmp_path):
        p = self._write_ics(tmp_path)
        result = ics_utils.parse_ics(str(p))
        assert result["comments"] == ""

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            ics_utils.parse_ics(str(tmp_path / "missing.ics"))

    def test_no_vevent_raises(self, tmp_path):
        p = tmp_path / "empty.ics"
        p.write_text("BEGIN:VCALENDAR\nEND:VCALENDAR\n", encoding="utf-8")
        with pytest.raises(ValueError):
            ics_utils.parse_ics(str(p))


# -----------------------------------------------------------------
# parse_ics_events (task #94): whole-calendar reading, full start/end
# datetimes, for correlate_calendar_recordings.py.
# -----------------------------------------------------------------

MULTI_EVENT_ICS = """BEGIN:VCALENDAR
VERSION:2.0
BEGIN:VEVENT
SUMMARY:Morning Standup
DTSTART:20260704T090000
DTEND:20260704T091500
END:VEVENT
BEGIN:VEVENT
SUMMARY:Customer Call
DTSTART:20260704T140000
DTEND:20260704T150000
DESCRIPTION:Discuss GC-MS quote.
END:VEVENT
END:VCALENDAR
"""


class TestParseIcsEvents:
    def test_reads_every_vevent_in_order(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(MULTI_EVENT_ICS, encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert [e["title"] for e in events] == ["Morning Standup", "Customer Call"]

    def test_start_and_end_are_datetimes(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(MULTI_EVENT_ICS, encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        first = events[0]
        assert first["start"] == datetime(2026, 7, 4, 9, 0, 0)
        assert first["end"] == datetime(2026, 7, 4, 9, 15, 0)
        assert first["all_day"] is False

    def test_description_becomes_comments(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(MULTI_EVENT_ICS, encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert events[1]["comments"] == "Discuss GC-MS quote."
        assert events[0]["comments"] == ""

    def test_missing_dtend_falls_back_to_dtstart(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(
            "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Point in time\n"
            "DTSTART:20260704T100000\nEND:VEVENT\nEND:VCALENDAR\n", encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert events[0]["start"] == datetime(2026, 7, 4, 10, 0, 0)
        assert events[0]["end"] == datetime(2026, 7, 4, 10, 0, 0)

    def test_all_day_event_has_no_time_component(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(
            "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Company Holiday\n"
            "DTSTART:20260704\nDTEND:20260705\nEND:VEVENT\nEND:VCALENDAR\n", encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert events[0]["all_day"] is True
        assert events[0]["start"] == datetime(2026, 7, 4, 0, 0, 0)
        assert events[0]["end"] == datetime(2026, 7, 5, 0, 0, 0)

    def test_utc_z_suffix_is_parsed_without_raising(self, tmp_path):
        # No TZID conversion is attempted (see module docstring) - just
        # confirm a "Z"-suffixed UTC value parses to *some* datetime
        # rather than being dropped or raising.
        p = tmp_path / "calendar.ics"
        p.write_text(
            "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:UTC Event\n"
            "DTSTART:20260704T120000Z\nDTEND:20260704T130000Z\n"
            "END:VEVENT\nEND:VCALENDAR\n", encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert isinstance(events[0]["start"], datetime)
        assert isinstance(events[0]["end"], datetime)
        assert events[0]["end"] > events[0]["start"]

    def test_no_vevent_returns_empty_list_not_raise(self, tmp_path):
        p = tmp_path / "empty.ics"
        p.write_text("BEGIN:VCALENDAR\nEND:VCALENDAR\n", encoding="utf-8")
        assert ics_utils.parse_ics_events(str(p)) == []

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            ics_utils.parse_ics_events(str(tmp_path / "missing.ics"))

    def test_malformed_dtstart_yields_none_start_and_end(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(
            "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Broken\n"
            "DTSTART:not-a-date\nEND:VEVENT\nEND:VCALENDAR\n", encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert events[0]["start"] is None
        assert events[0]["end"] is None

    def test_unterminated_vevent_at_eof_is_not_parsed(self, tmp_path):
        p = tmp_path / "calendar.ics"
        p.write_text(
            "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Complete\n"
            "DTSTART:20260704T090000\nEND:VEVENT\n"
            "BEGIN:VEVENT\nSUMMARY:Dangling\nDTSTART:20260704T100000\n",
            encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert [e["title"] for e in events] == ["Complete"]


# -----------------------------------------------------------------
# Attendees and Agenda (task #95): ATTENDEE lines and a best-effort
# "Agenda:" section inside DESCRIPTION, read by both parse_ics (single
# event) and parse_ics_events (whole calendar).
# -----------------------------------------------------------------

ATTENDEE_ICS = (
    "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Customer Call\n"
    "DTSTART:20260704T140000\nDTEND:20260704T150000\n"
    "ORGANIZER;CN=Joerg Riener:mailto:joergriener@gmail.com\n"
    "ATTENDEE;CN=Jane Doe;ROLE=REQ-PARTICIPANT:mailto:jane@example.com\n"
    "ATTENDEE;CN=\"Smith, John\";ROLE=REQ-PARTICIPANT:mailto:john@example.com\n"
    "ATTENDEE:mailto:noname@example.com\n"
    "DESCRIPTION:Agenda:\\n1. Intro\\n2. Demo\\n\\nJoin link: https://example.com\n"
    "END:VEVENT\nEND:VCALENDAR\n"
)


class TestParseIcsAttendeesAndAgenda:
    def test_parse_ics_returns_attendees(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(ATTENDEE_ICS, encoding="utf-8")
        result = ics_utils.parse_ics(str(p))
        names = [a["name"] for a in result["attendees"]]
        assert names == ["Jane Doe", "Smith, John", "noname@example.com"]

    def test_organizer_is_not_included_as_attendee(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(ATTENDEE_ICS, encoding="utf-8")
        result = ics_utils.parse_ics(str(p))
        names = [a["name"] for a in result["attendees"]]
        assert "Joerg Riener" not in names

    def test_attendee_email_captured(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(ATTENDEE_ICS, encoding="utf-8")
        result = ics_utils.parse_ics(str(p))
        assert result["attendees"][0]["email"] == "jane@example.com"

    def test_no_attendee_lines_yields_empty_list(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(MULTI_EVENT_ICS, encoding="utf-8")
        result = ics_utils.parse_ics(str(p))
        assert result["attendees"] == []

    def test_agenda_extracted_from_description(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(ATTENDEE_ICS, encoding="utf-8")
        result = ics_utils.parse_ics(str(p))
        assert result["agenda"] == "1. Intro 2. Demo"

    def test_agenda_blank_when_no_label_present(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(MULTI_EVENT_ICS, encoding="utf-8")
        result = ics_utils.parse_ics(str(p))
        assert result["agenda"] == ""

    def test_parse_ics_events_carries_attendees_and_agenda_per_event(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(ATTENDEE_ICS, encoding="utf-8")
        events = ics_utils.parse_ics_events(str(p))
        assert [a["name"] for a in events[0]["attendees"]] == \
            ["Jane Doe", "Smith, John", "noname@example.com"]
        assert events[0]["agenda"] == "1. Intro 2. Demo"


class TestFormatAttendeesAndAttendeeNames:
    def test_format_attendees_joins_name_and_email(self):
        attendees = [{"name": "Jane Doe", "email": "jane@example.com"},
                     {"name": "John", "email": ""}]
        assert ics_utils.format_attendees(attendees) == \
            "Jane Doe <jane@example.com>; John"

    def test_format_attendees_empty_list(self):
        assert ics_utils.format_attendees([]) == ""

    def test_attendee_names_falls_back_to_email(self):
        attendees = [{"name": "", "email": "jane@example.com"}, {"name": "John", "email": ""}]
        assert ics_utils.attendee_names(attendees) == ["jane@example.com", "John"]

    def test_attendee_names_empty_list(self):
        assert ics_utils.attendee_names([]) == []


# -----------------------------------------------------------------
# parse_labeled_text (shared by .txt and .docx)
# -----------------------------------------------------------------

class TestParseLabeledText:
    def test_all_three_fields(self):
        text = "Title: Quarterly Review\nDate: 2026-07-04\nComments: Bring laptop.\n"
        result = ics_utils.parse_labeled_text(text)
        assert result == {"title": "Quarterly Review", "date": "2026-07-04",
                           "comments": "Bring laptop.", "agenda": "", "attendees": []}

    def test_comments_span_multiple_lines_until_next_label(self):
        text = (
            "Title: Kickoff\n"
            "Comments: First line of notes.\n"
            "Second line of notes.\n"
            "Third line.\n"
            "Date: 2026-07-04\n"
        )
        result = ics_utils.parse_labeled_text(text)
        assert result["comments"] == "First line of notes.\nSecond line of notes.\nThird line."
        assert result["date"] == "2026-07-04"

    def test_synonyms_recognized(self):
        text = "Subject: Board Meeting\nNotes: Confidential agenda.\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["title"] == "Board Meeting"
        assert result["comments"] == "Confidential agenda."

    def test_case_insensitive_labels(self):
        text = "TITLE: Standup\ncomments: quick sync\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["title"] == "Standup"
        assert result["comments"] == "quick sync"

    def test_ddmmyyyy_date_normalized_to_iso(self):
        text = "Date: 04.07.2026\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["date"] == "2026-07-04"

    def test_unrecognized_date_format_passed_through(self):
        text = "Date: next Tuesday\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["date"] == "next Tuesday"

    def test_lines_before_first_label_are_ignored(self):
        text = "Some preamble text.\nTitle: Real Title\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["title"] == "Real Title"

    def test_no_labels_returns_all_blank(self):
        text = "Just some random notes with no structure.\n"
        result = ics_utils.parse_labeled_text(text)
        assert result == {"title": "", "date": "", "comments": "", "agenda": "", "attendees": []}

    def test_blank_lines_between_comment_lines_are_skipped_not_joined_as_empty(self):
        text = "Comments: Line one.\n\nLine two.\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["comments"] == "Line one.\nLine two."

    def test_agenda_label_recognized(self):
        text = "Title: Kickoff\nAgenda: 1. Intro\n2. Demo\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["agenda"] == "1. Intro\n2. Demo"

    def test_invitees_one_per_line(self):
        text = "Invitees:\nJane Doe\nJohn Smith\n"
        result = ics_utils.parse_labeled_text(text)
        assert [a["name"] for a in result["attendees"]] == ["Jane Doe", "John Smith"]

    def test_invitees_comma_separated_on_one_line(self):
        text = "Invitees: Jane Doe, John Smith\n"
        result = ics_utils.parse_labeled_text(text)
        assert [a["name"] for a in result["attendees"]] == ["Jane Doe", "John Smith"]

    def test_invitees_with_email_in_angle_brackets(self):
        text = "Invitees: Jane Doe <jane@example.com>\n"
        result = ics_utils.parse_labeled_text(text)
        assert result["attendees"] == [{"name": "Jane Doe", "email": "jane@example.com"}]

    def test_attendees_and_participants_are_synonyms_for_invitees(self):
        assert ics_utils.parse_labeled_text("Attendees: Jane\n")["attendees"] == [
            {"name": "Jane", "email": ""}]
        assert ics_utils.parse_labeled_text("Participants: Jane\n")["attendees"] == [
            {"name": "Jane", "email": ""}]

    def test_no_invitees_label_yields_empty_list(self):
        result = ics_utils.parse_labeled_text("Title: Kickoff\n")
        assert result["attendees"] == []


# -----------------------------------------------------------------
# parse_txt
# -----------------------------------------------------------------

class TestParseTxt:
    def test_reads_labeled_txt_file(self, tmp_path):
        p = tmp_path / "meeting.txt"
        p.write_text("Title: Sales Call\nDate: 2026-07-04\nComments: Discuss pricing.\n",
                      encoding="utf-8")
        result = ics_utils.parse_txt(str(p))
        assert result["title"] == "Sales Call"
        assert result["comments"] == "Discuss pricing."

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            ics_utils.parse_txt(str(tmp_path / "missing.txt"))

    def test_no_recognized_labels_falls_back_to_whole_file_as_comments(self, tmp_path):
        """Task: YouTube description.txt import - a raw description has no
        Title:/Date:/Comments: labels at all, so parse_txt now uses the
        whole file as comments instead of refusing it outright (a
        properly labeled file, tested above, is unaffected)."""
        p = tmp_path / "plain.txt"
        p.write_text("Just some unrelated text.\nSecond line too.\n", encoding="utf-8")
        result = ics_utils.parse_txt(str(p))
        assert result["title"] == ""
        assert result["comments"] == "Just some unrelated text.\nSecond line too."

    def test_empty_file_still_raises(self, tmp_path):
        p = tmp_path / "empty.txt"
        p.write_text("", encoding="utf-8")
        with pytest.raises(ValueError):
            ics_utils.parse_txt(str(p))

    def test_whitespace_only_file_still_raises(self, tmp_path):
        p = tmp_path / "blank.txt"
        p.write_text("   \n\n   \n", encoding="utf-8")
        with pytest.raises(ValueError):
            ics_utils.parse_txt(str(p))


# -----------------------------------------------------------------
# parse_docx
# -----------------------------------------------------------------

class TestParseDocx:
    def _write_docx(self, tmp_path, lines):
        from docx import Document
        doc = Document()
        for line in lines:
            doc.add_paragraph(line)
        p = tmp_path / "meeting.docx"
        doc.save(str(p))
        return p

    def test_reads_labeled_docx_paragraphs(self, tmp_path):
        p = self._write_docx(tmp_path, [
            "Title: Product Demo", "Date: 2026-07-04",
            "Comments: Show the new dashboard first.",
        ])
        result = ics_utils.parse_docx(str(p))
        assert result["title"] == "Product Demo"
        assert result["date"] == "2026-07-04"
        assert result["comments"] == "Show the new dashboard first."

    def test_comments_across_multiple_paragraphs(self, tmp_path):
        p = self._write_docx(tmp_path, [
            "Comments: Paragraph one.", "Paragraph two.", "Title: Wrap-up",
        ])
        result = ics_utils.parse_docx(str(p))
        assert result["comments"] == "Paragraph one.\nParagraph two."
        assert result["title"] == "Wrap-up"

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            ics_utils.parse_docx(str(tmp_path / "missing.docx"))

    def test_no_recognized_labels_raises(self, tmp_path):
        p = self._write_docx(tmp_path, ["Just some unrelated paragraph."])
        with pytest.raises(ValueError):
            ics_utils.parse_docx(str(p))


# -----------------------------------------------------------------
# parse_meeting_info_file dispatcher
# -----------------------------------------------------------------

class TestParseMeetingInfoFile:
    def test_dispatches_ics(self, tmp_path):
        p = tmp_path / "invite.ics"
        p.write_text(
            "BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Call\nDTSTART:20260704T090000\n"
            "END:VEVENT\nEND:VCALENDAR\n", encoding="utf-8")
        result = ics_utils.parse_meeting_info_file(str(p))
        assert result["title"] == "Call"

    def test_dispatches_txt(self, tmp_path):
        p = tmp_path / "notes.txt"
        p.write_text("Title: Text Dispatch\n", encoding="utf-8")
        result = ics_utils.parse_meeting_info_file(str(p))
        assert result["title"] == "Text Dispatch"

    def test_dispatches_docx(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Title: Docx Dispatch")
        p = tmp_path / "notes.docx"
        doc.save(str(p))
        result = ics_utils.parse_meeting_info_file(str(p))
        assert result["title"] == "Docx Dispatch"

    def test_unsupported_extension_raises(self, tmp_path):
        p = tmp_path / "notes.pdf"
        p.write_text("irrelevant", encoding="utf-8")
        with pytest.raises(ValueError):
            ics_utils.parse_meeting_info_file(str(p))
