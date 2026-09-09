# =============================================================
#  Transkription_Notes_Pipeline - tests/test_doc_translate.py
#  Unit tests for doc_translate.py's per-format readers/writers.
#  translate_fn is a deterministic fake throughout (no live LLM call);
#  these tests check document STRUCTURE (paragraphs, runs, tables,
#  speaker notes, formatting survival) is preserved correctly, not
#  translation quality.
# =============================================================

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest
import doc_translate


def fake_translate(text: str) -> str:
    """Deterministic stand-in for an LLM call: wraps input in markers so
    tests can assert exactly what doc_translate.py sent for translation."""
    return f"[[{text}]]"


class TestTxt:
    def test_paragraphs_translated_and_rejoined(self, tmp_path):
        src = tmp_path / "in.txt"
        src.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")
        dest = tmp_path / "out.txt"
        doc_translate.translate_txt(src, dest, fake_translate)
        result = dest.read_text(encoding="utf-8")
        assert result == "[[First paragraph.]]\n\n[[Second paragraph.]]"

    def test_blank_lines_dropped_not_translated(self, tmp_path):
        src = tmp_path / "in.txt"
        src.write_text("Only paragraph.\n\n\n\n", encoding="utf-8")
        dest = tmp_path / "out.txt"
        doc_translate.translate_txt(src, dest, fake_translate)
        assert dest.read_text(encoding="utf-8") == "[[Only paragraph.]]"


class TestDocx:
    def test_paragraph_text_translated_in_place(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello world.")
        src = tmp_path / "in.docx"
        doc.save(str(src))

        dest = tmp_path / "out.docx"
        doc_translate.translate_docx(src, dest, fake_translate)

        result = Document(str(dest))
        texts = [p.text for p in result.paragraphs if p.text.strip()]
        assert texts == ["[[Hello world.]]"]

    def test_multi_run_paragraph_keeps_first_runs_formatting(self, tmp_path):
        from docx import Document
        doc = Document()
        p = doc.add_paragraph()
        run1 = p.add_run("Please ")
        run1.bold = True
        run2 = p.add_run("confirm by Friday.")
        run2.italic = True
        src = tmp_path / "in.docx"
        doc.save(str(src))

        dest = tmp_path / "out.docx"
        doc_translate.translate_docx(src, dest, fake_translate)

        result = Document(str(dest))
        para = result.paragraphs[0]
        assert para.text == "[[Please confirm by Friday.]]"
        assert para.runs[0].bold is True
        assert para.runs[1].text == ""

    def test_table_cells_translated(self, tmp_path):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        src = tmp_path / "in.docx"
        doc.save(str(src))

        dest = tmp_path / "out.docx"
        doc_translate.translate_docx(src, dest, fake_translate)

        result = Document(str(dest))
        cells_text = [c.text for row in result.tables[0].rows for c in row.cells]
        assert cells_text == ["[[Name]]", "[[Value]]"]

    def test_blank_paragraph_left_untouched(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("")
        doc.add_paragraph("Real text.")
        src = tmp_path / "in.docx"
        doc.save(str(src))

        dest = tmp_path / "out.docx"
        doc_translate.translate_docx(src, dest, fake_translate)

        result = Document(str(dest))
        texts = [p.text for p in result.paragraphs]
        assert texts[0] == ""
        assert texts[1] == "[[Real text.]]"

    def test_repeated_paragraph_translated_only_once(self, tmp_path):
        from docx import Document
        calls = []

        def counting_translate(text):
            calls.append(text)
            return f"[[{text}]]"

        doc = Document()
        doc.add_paragraph("Confidential")
        doc.add_paragraph("Confidential")
        src = tmp_path / "in.docx"
        doc.save(str(src))

        dest = tmp_path / "out.docx"
        doc_translate.translate_docx(src, dest, counting_translate)

        assert calls == ["Confidential"]  # cached on the second occurrence


class TestPptx:
    def test_slide_text_and_speaker_notes_translated(self, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "My Title"
        body = slide.placeholders[1]
        body.text = "Bullet point one."
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = "Remember to mention Q3 numbers."
        src = tmp_path / "in.pptx"
        prs.save(str(src))

        dest = tmp_path / "out.pptx"
        doc_translate.translate_pptx(src, dest, fake_translate)

        result = Presentation(str(dest))
        result_slide = result.slides[0]
        assert result_slide.shapes.title.text == "[[My Title]]"
        assert result_slide.placeholders[1].text == "[[Bullet point one.]]"
        assert result_slide.notes_slide.notes_text_frame.text == \
            "[[Remember to mention Q3 numbers.]]"

    def test_table_on_slide_translated(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        graphic_frame = slide.shapes.add_table(
            rows=1, cols=2, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(1))
        graphic_frame.table.cell(0, 0).text = "Metric"
        graphic_frame.table.cell(0, 1).text = "Result"
        src = tmp_path / "in.pptx"
        prs.save(str(src))

        dest = tmp_path / "out.pptx"
        doc_translate.translate_pptx(src, dest, fake_translate)

        result = Presentation(str(dest))
        table = result.slides[0].shapes[0].table
        assert table.cell(0, 0).text == "[[Metric]]"
        assert table.cell(0, 1).text == "[[Result]]"


class TestHtml:
    SAMPLE_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>MEETING NOTES: call.mp4</title>
<style>
  .header { color: #1F5C99; font-size: 28px; }
  h1::before { content: "Section: "; }
</style>
</head>
<body>
  <div class="header">
    <h1>Meeting Notes</h1>
    <span class="filename">call.mp4</span>
  </div>
  <div class="content">
    <h2>Main Topics</h2>
    <ul>
      <li>First key point.</li>
      <li>Second key point.</li>
    </ul>
    <img src="snapshots/slide001.png" alt="Title slide">
  </div>
  <div class="footer">Generated by Transkription_Notes_Pipeline</div>
</body>
</html>"""

    def test_style_and_script_untouched(self, tmp_path):
        src = tmp_path / "in.html"
        src.write_text(self.SAMPLE_HTML, encoding="utf-8")
        dest = tmp_path / "out.html"
        doc_translate.translate_html(src, dest, fake_translate)
        result = dest.read_text(encoding="utf-8")
        assert "color: #1F5C99; font-size: 28px;" in result
        assert 'content: "Section: ";' in result

    def test_visible_text_translated(self, tmp_path):
        src = tmp_path / "in.html"
        src.write_text(self.SAMPLE_HTML, encoding="utf-8")
        dest = tmp_path / "out.html"
        doc_translate.translate_html(src, dest, fake_translate)
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(dest.read_text(encoding="utf-8"), "html.parser")
        assert soup.h1.text == "[[Meeting Notes]]"
        assert soup.h2.text == "[[Main Topics]]"
        lis = soup.find_all("li")
        assert [li.text for li in lis] == ["[[First key point.]]", "[[Second key point.]]"]
        assert soup.title.text == "[[MEETING NOTES: call.mp4]]"

    def test_alt_attribute_translated_src_untouched(self, tmp_path):
        src = tmp_path / "in.html"
        src.write_text(self.SAMPLE_HTML, encoding="utf-8")
        dest = tmp_path / "out.html"
        doc_translate.translate_html(src, dest, fake_translate)
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(dest.read_text(encoding="utf-8"), "html.parser")
        img = soup.find("img")
        assert img["alt"] == "[[Title slide]]"
        assert img["src"] == "snapshots/slide001.png"

    def test_class_and_tag_structure_preserved(self, tmp_path):
        src = tmp_path / "in.html"
        src.write_text(self.SAMPLE_HTML, encoding="utf-8")
        dest = tmp_path / "out.html"
        doc_translate.translate_html(src, dest, fake_translate)
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(dest.read_text(encoding="utf-8"), "html.parser")
        assert soup.find("div", class_="header") is not None
        assert soup.find("div", class_="footer").text == "[[Generated by Transkription_Notes_Pipeline]]"

    def test_output_suffix_html_and_htm_unchanged(self):
        assert doc_translate.output_suffix_for(Path("a.html")) == ".html"
        assert doc_translate.output_suffix_for(Path("a.htm")) == ".htm"

    def test_translate_file_dispatches_html(self, tmp_path):
        src = tmp_path / "report.html"
        src.write_text(self.SAMPLE_HTML, encoding="utf-8")
        dest = doc_translate.translate_file(str(src), "de", fake_translate)
        assert dest == tmp_path / "report_de.html"
        assert "[[Meeting Notes]]" in dest.read_text(encoding="utf-8")


class TestImageOcrLanguage:
    """
    OCR language plumbing - see doc_translate._ocr_image_text's docstring
    for the actual bug this guards against (Tesseract silently misreading
    a non-Latin script under the wrong/default language pack, producing
    garbled text that a capable LLM then "translates" into fluent-sounding
    nonsense). pytesseract.image_to_string is monkeypatched throughout -
    no real OCR engine involved, this only checks the right lang value
    reaches Tesseract and that a missing-language-pack error degrades
    gracefully instead of failing the whole file.
    """

    def _make_image(self, tmp_path) -> Path:
        from PIL import Image
        img = Image.new("RGB", (20, 20), "white")
        src = tmp_path / "shot.png"
        img.save(src)
        return src

    def test_configured_lang_passed_to_tesseract(self, tmp_path, monkeypatch):
        import pytesseract
        captured = {}

        def fake_image_to_string(img, lang=None):
            captured["lang"] = lang
            return "Hallo Welt"

        monkeypatch.setattr(pytesseract, "image_to_string", fake_image_to_string)
        src = self._make_image(tmp_path)
        dest = tmp_path / "shot_de.docx"
        doc_translate.translate_image_to_docx(src, dest, fake_translate, ocr_lang="deu")
        assert captured["lang"] == "deu"

    def test_default_lang_is_eng_plus_deu(self, tmp_path, monkeypatch):
        import pytesseract
        captured = {}

        def fake_image_to_string(img, lang=None):
            captured["lang"] = lang
            return "text"

        monkeypatch.setattr(pytesseract, "image_to_string", fake_image_to_string)
        src = self._make_image(tmp_path)
        dest = tmp_path / "shot_en.docx"
        doc_translate.translate_image_to_docx(src, dest, fake_translate)
        assert captured["lang"] == "eng+deu"

    def test_missing_language_pack_falls_back_without_crashing(self, tmp_path, monkeypatch):
        import pytesseract
        calls = []

        def flaky_image_to_string(img, lang=None):
            calls.append(lang)
            if lang == "chi_sim":
                raise pytesseract.TesseractError(1, "Error opening data file chi_sim.traineddata")
            return "Fallback text recognised without the missing pack"

        monkeypatch.setattr(pytesseract, "image_to_string", flaky_image_to_string)
        src = self._make_image(tmp_path)
        dest = tmp_path / "shot_en.docx"
        doc_translate.translate_image_to_docx(src, dest, fake_translate, ocr_lang="chi_sim")
        # First call with the requested (missing) pack, then a bare retry
        # with no lang argument at all (Tesseract's own default) - never
        # raises out to the caller, the file still gets processed.
        assert calls == ["chi_sim", None]
        assert dest.exists()

    def test_image_converted_to_grayscale_before_ocr(self, tmp_path, monkeypatch):
        captured = {}

        def fake_image_to_string(img, lang=None):
            captured["mode"] = img.mode
            return "text"

        import pytesseract
        monkeypatch.setattr(pytesseract, "image_to_string", fake_image_to_string)
        src = self._make_image(tmp_path)  # created as RGB
        dest = tmp_path / "shot_en.docx"
        doc_translate.translate_image_to_docx(src, dest, fake_translate)
        assert captured["mode"] == "L"


class TestDispatch:
    def test_output_suffix_docx_pptx_txt_unchanged(self):
        assert doc_translate.output_suffix_for(Path("a.docx")) == ".docx"
        assert doc_translate.output_suffix_for(Path("a.pptx")) == ".pptx"
        assert doc_translate.output_suffix_for(Path("a.txt")) == ".txt"

    def test_output_suffix_pdf_and_images_become_docx(self):
        assert doc_translate.output_suffix_for(Path("a.pdf")) == ".docx"
        assert doc_translate.output_suffix_for(Path("a.png")) == ".docx"
        assert doc_translate.output_suffix_for(Path("a.jpg")) == ".docx"

    def test_legacy_doc_raises_clear_error(self, tmp_path):
        src = tmp_path / "old.doc"
        src.write_bytes(b"not a real doc file")
        with pytest.raises(ValueError, match=r"\.docx"):
            doc_translate.translate_file(str(src), "de", fake_translate)

    def test_legacy_ppt_raises_clear_error(self, tmp_path):
        src = tmp_path / "old.ppt"
        src.write_bytes(b"not a real ppt file")
        with pytest.raises(ValueError, match=r"\.pptx"):
            doc_translate.translate_file(str(src), "de", fake_translate)

    def test_unsupported_extension_raises(self, tmp_path):
        src = tmp_path / "a.xyz"
        src.write_text("hi", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported file type"):
            doc_translate.translate_file(str(src), "de", fake_translate)

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            doc_translate.translate_file(str(tmp_path / "nope.txt"), "de", fake_translate)

    def test_translate_file_writes_expected_name_next_to_source(self, tmp_path):
        src = tmp_path / "meeting.txt"
        src.write_text("Hello.", encoding="utf-8")
        dest = doc_translate.translate_file(str(src), "de", fake_translate)
        assert dest == tmp_path / "meeting_de.txt"
        assert dest.exists()

    def test_translate_file_respects_output_dir_override(self, tmp_path):
        src_dir = tmp_path / "source"
        src_dir.mkdir()
        out_dir = tmp_path / "out"
        src = src_dir / "meeting.txt"
        src.write_text("Hello.", encoding="utf-8")
        dest = doc_translate.translate_file(str(src), "de", fake_translate,
                                            output_dir_override=str(out_dir))
        assert dest == out_dir / "meeting_de.txt"
        assert dest.exists()


class TestPdf:
    def test_no_text_layer_raises_clear_error(self, tmp_path):
        # A minimal valid empty PDF has no extractable text.
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        src = tmp_path / "blank.pdf"
        with open(src, "wb") as f:
            writer.write(f)

        dest = tmp_path / "blank_de.docx"
        with pytest.raises(ValueError, match="No extractable text"):
            doc_translate.translate_pdf_to_docx(src, dest, fake_translate)
