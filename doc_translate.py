# =============================================================
#  Transkription_Notes_Pipeline - doc_translate.py
#  File-format I/O for the "Translate" tab/CLI: reads text out of
#  txt/docx/pptx/pdf/image files, translates it via a caller-supplied
#  translate_fn(text) -> str|None (see translator.translate_text; wired
#  in by run_pipeline.py so this module has no LLM/network code at
#  all), and writes the result back in a sensible output format.
#
#  Design: translate paragraph-by-paragraph (docx paragraphs, pptx
#  shape/table/notes paragraphs, txt/pdf/image blank-line-separated
#  paragraphs) rather than as one giant blob per file. This keeps
#  document structure (headings, bullet lists, slide-by-slide notes)
#  intact and keeps each LLM call small, at the cost of more calls for
#  very large documents - repeated identical paragraphs (a recurring
#  footer, "Confidential", etc.) are cached per run so they are only
#  translated once.
#
#  In-place formatting (docx/pptx): a paragraph's full text is
#  translated as ONE unit, then written back into its FIRST run (which
#  keeps that run's font/size/color/bold/etc.), and every other run in
#  the same paragraph is cleared. This is the standard practical
#  trade-off used by most docx/pptx translation tools: perfect
#  paragraph-level formatting, imperfect run-level formatting (a
#  paragraph that was half-bold loses that distinction mid-sentence).
#  Translating run-by-run instead would avoid that but breaks grammar
#  whenever a sentence spans multiple runs, which is far more common
#  and far more noticeable than a lost mid-paragraph style change.
#  Images, table cell structure, and slide layout/positions are
#  untouched either way - only paragraph text is rewritten.
#
#  Supported inputs: .txt, .docx, .pptx, .html/.htm (this project's own
#  generated notes/slide-report HTML included), .pdf (-> translated
#  .docx), .jpg/.jpeg/.png/.bmp/.tiff (-> translated .docx, via
#  Tesseract OCR). NOT supported: legacy .doc/.ppt (binary, pre-2007
#  formats) - raises a clear error asking the user to save-as .docx/
#  .pptx first (python-docx/python-pptx, like this project's existing
#  python-docx notes export, only read the modern XML-based formats).
# =============================================================

import logging
from pathlib import Path

log = logging.getLogger(__name__)

# Formats this module can read directly.
TRANSLATABLE_EXTENSIONS = {
    ".txt", ".docx", ".pptx", ".pdf", ".html", ".htm",
    ".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif",
}

# Legacy binary formats python-docx/python-pptx cannot open at all.
UNSUPPORTED_LEGACY_EXTENSIONS = {".doc", ".ppt"}

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif"}


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

def _split_paragraphs(text: str) -> list[str]:
    """Blank-line-separated paragraphs, used for txt/pdf/image extracted
    text (which has no native paragraph objects to iterate over, unlike
    docx/pptx)."""
    paragraphs = [p.strip() for p in text.replace("\r\n", "\n").split("\n\n")]
    return [p for p in paragraphs if p]


def _translate_unit(text: str, translate_fn, cache: dict) -> str:
    """
    translate_fn(text) -> str|None wrapper with a small in-memory cache
    keyed on the exact source string, scoped to one file's translation
    run: many docx/pptx paragraphs repeat verbatim (recurring footers,
    "Confidential", a title repeated on every slide) - avoid re-sending
    the identical string to the LLM more than once. Falls back to the
    original text if translate_fn returns None (failure) rather than
    dropping the content silently.
    """
    if not text or not text.strip():
        return text or ""
    if text in cache:
        return cache[text]
    result = translate_fn(text)
    cache[text] = result if result is not None else text
    return cache[text]


# ---------------------------------------------------------------------------
# .txt
# ---------------------------------------------------------------------------

def translate_txt(src_path: Path, dest_path: Path, translate_fn) -> None:
    text = src_path.read_text(encoding="utf-8", errors="replace")
    cache: dict = {}
    paragraphs = _split_paragraphs(text)
    translated = [_translate_unit(p, translate_fn, cache) for p in paragraphs]
    dest_path.write_text("\n\n".join(translated), encoding="utf-8")


# ---------------------------------------------------------------------------
# .docx  (also used for the paragraph-in-place helper shared with .pptx)
# ---------------------------------------------------------------------------

def _translate_paragraph_runs(paragraph, translate_fn, cache: dict) -> None:
    """
    Translate one python-docx/python-pptx paragraph object in place:
    full-paragraph text -> first run keeps it, remaining runs cleared.
    No-op for a blank/whitespace-only paragraph or one with no runs at
    all (both libraries expose the same .text getter and .runs list
    with a writable run.text, so this works unchanged for either).
    """
    text = paragraph.text
    if not text or not text.strip():
        return
    if not paragraph.runs:
        return
    translated = _translate_unit(text, translate_fn, cache)
    paragraph.runs[0].text = translated
    for run in paragraph.runs[1:]:
        run.text = ""


def translate_docx(src_path: Path, dest_path: Path, translate_fn) -> None:
    from docx import Document
    doc = Document(str(src_path))
    cache: dict = {}
    for paragraph in doc.paragraphs:
        _translate_paragraph_runs(paragraph, translate_fn, cache)
    # Table cells (top-level tables only; a table nested inside another
    # table's cell is a rare enough document structure that it is left
    # untranslated in this first version rather than adding recursion).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _translate_paragraph_runs(paragraph, translate_fn, cache)
    doc.save(str(dest_path))


# ---------------------------------------------------------------------------
# .pptx  (slide text, tables, AND speaker notes)
# ---------------------------------------------------------------------------

def _translate_text_frame(text_frame, translate_fn, cache: dict) -> None:
    for paragraph in text_frame.paragraphs:
        _translate_paragraph_runs(paragraph, translate_fn, cache)


def translate_pptx(src_path: Path, dest_path: Path, translate_fn) -> None:
    from pptx import Presentation
    prs = Presentation(str(src_path))
    cache: dict = {}
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                _translate_text_frame(shape.text_frame, translate_fn, cache)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _translate_text_frame(cell.text_frame, translate_fn, cache)
        # Speaker notes (explicitly requested): translated the same way
        # as any other text frame, so they end up in the same output
        # language as the slide content itself.
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf is not None:
                _translate_text_frame(notes_tf, translate_fn, cache)
    prs.save(str(dest_path))


# ---------------------------------------------------------------------------
# .html / .htm  (both notes.py's save_notes_html AND reporter.py's slide
# report HTML - both are plain HTML with an embedded <style> block, no
# <script>). Kept in its own format (like docx/pptx/txt), not converted
# to .docx: the whole point of translating a generated report is usually
# to still open it as a report (styling, embedded slide images, layout
# all intact), not to get a plain Word document.
#
# Approach: parse with BeautifulSoup, translate only actual text nodes
# (skipping <script>/<style>, whose content is code/CSS, not natural-
# language text) plus a couple of human-readable attributes (alt, title
# tooltip). Every tag, class, id, inline style, and src/href URL is left
# completely untouched - only visible text changes.
# ---------------------------------------------------------------------------

def translate_html(src_path: Path, dest_path: Path, translate_fn) -> None:
    from bs4 import BeautifulSoup, Comment

    html_content = src_path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html_content, "html.parser")
    cache: dict = {}

    skip_tags = {"script", "style"}
    for node in soup.find_all(string=True):
        if isinstance(node, Comment):
            continue
        if node.parent is not None and node.parent.name in skip_tags:
            continue
        text = str(node)
        if not text.strip():
            continue
        # Preserve the original leading/trailing whitespace (template
        # indentation/newlines) - only the stripped core text is sent
        # for translation, so re-serializing the tree doesn't collapse
        # the source template's formatting.
        leading = text[:len(text) - len(text.lstrip())]
        trailing = text[len(text.rstrip()):]
        core = text.strip()
        translated = _translate_unit(core, translate_fn, cache)
        node.replace_with(leading + translated + trailing)

    for attr in ("alt", "title"):
        for tag in soup.find_all(attrs={attr: True}):
            value = tag.get(attr, "")
            if value.strip():
                tag[attr] = _translate_unit(value.strip(), translate_fn, cache)

    dest_path.write_text(str(soup), encoding="utf-8")


# ---------------------------------------------------------------------------
# .pdf -> translated .docx
#
# PDF has no editable layout to write a translation back into (unlike
# docx/pptx, which are XML documents with a real paragraph/run model).
# Per project decision, extracted text is translated and written into a
# clean new Word document instead of attempting to reconstruct the PDF's
# original visual layout, which pypdf (or any pure-Python PDF library)
# cannot do reliably for arbitrary PDFs anyway.
# ---------------------------------------------------------------------------

def _extract_pdf_text(src_path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(src_path))
    pages_text = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(pages_text)


def _write_translated_docx(paragraphs: list[str], src_name: str, dest_path: Path,
                           translate_fn) -> None:
    from docx import Document
    cache: dict = {}
    doc = Document()
    doc.add_heading(f"Translated: {src_name}", level=1)
    for para in paragraphs:
        translated = _translate_unit(para, translate_fn, cache)
        doc.add_paragraph(translated)
    doc.save(str(dest_path))


def translate_pdf_to_docx(src_path: Path, dest_path: Path, translate_fn) -> None:
    text = _extract_pdf_text(src_path)
    if not text.strip():
        raise ValueError(
            f"No extractable text found in {src_path.name} (likely a scanned/"
            "image-only PDF with no text layer - try exporting its pages as "
            "images and use the image/OCR path instead)."
        )
    paragraphs = _split_paragraphs(text)
    _write_translated_docx(paragraphs, src_path.name, dest_path, translate_fn)


# ---------------------------------------------------------------------------
# Images -> translated .docx, via Tesseract OCR
#
# Requires the pytesseract Python package AND the Tesseract-OCR binary
# itself installed and on PATH (https://github.com/tesseract-ocr/tesseract;
# this project's README already notes a Tesseract-OCR install is common
# on machines also running DocsSorter/Reisekosten_Scans).
#
# lang matters a lot here: Tesseract's default ("eng" unless the caller
# says otherwise) reads whatever glyphs it sees through an English
# letter-shape model. Fed a non-Latin script (Chinese, Japanese, etc.)
# this does not fail loudly - it happily emits confident-looking but
# completely wrong Latin-lookalike text ("AMTARIER EAI A Se = ae GB" for
# a Chinese slide, observed 2026-07-18). translate_text then faithfully
# "translates" that garbage, and a capable LLM will often smooth it into
# fluent-sounding but entirely invented sentences rather than fail
# visibly - the end result looks like a translation bug but the actual
# corruption already happened at the OCR step, before translation was
# ever involved. Always set lang to match the slide's real script (see
# config.TESSERACT_LANG_MAP / translate_ocr_language) rather than
# relying on the default.
# ---------------------------------------------------------------------------

def _ocr_image_text(src_path: Path, lang: str = "eng+deu") -> str:
    import pytesseract
    from PIL import Image
    with Image.open(src_path) as img:
        # Grayscale is a safe, standard Tesseract accuracy improvement
        # for screenshots/slides (as opposed to clean scanned pages);
        # does not affect which language pack is used.
        img = img.convert("L")
        try:
            return pytesseract.image_to_string(img, lang=lang)
        except pytesseract.TesseractError as exc:
            log.warning(
                "Tesseract language '%s' unavailable for %s (%s) - falling back to "
                "Tesseract's own default. Run `tesseract --list-langs` to see installed "
                "packs; install the missing one from "
                "https://github.com/tesseract-ocr/tessdata if OCR quality looks wrong.",
                lang, src_path.name, exc,
            )
            return pytesseract.image_to_string(img)


def translate_image_to_docx(src_path: Path, dest_path: Path, translate_fn,
                            ocr_lang: str = "eng+deu") -> None:
    try:
        text = _ocr_image_text(src_path, lang=ocr_lang)
    except Exception as exc:
        raise RuntimeError(
            f"OCR failed for {src_path.name}: {exc}. Requires the pytesseract "
            "Python package AND the Tesseract-OCR binary installed and on PATH "
            "(https://github.com/tesseract-ocr/tesseract)."
        ) from exc
    if not text.strip():
        raise ValueError(f"No text detected in {src_path.name} (OCR found nothing).")
    paragraphs = _split_paragraphs(text)
    _write_translated_docx(paragraphs, src_path.name, dest_path, translate_fn)


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

def output_suffix_for(src_path: Path) -> str:
    """Extension the translated output will be written with. docx/pptx/
    txt/html keep their own format; pdf/images become .docx (see module
    docstring for why)."""
    ext = src_path.suffix.lower()
    if ext in (".docx", ".pptx", ".txt", ".html", ".htm"):
        return ext
    if ext == ".pdf" or ext in IMAGE_EXTENSIONS:
        return ".docx"
    raise ValueError(f"Unsupported file type: {ext}")


def translate_file(src_path: str, target_language: str, translate_fn,
                   output_dir_override: str = "", ocr_lang: str = "eng+deu") -> Path:
    """
    Translate one file, dispatching on extension. translate_fn(text) ->
    str|None does the actual LLM call (see translator.translate_text);
    this module stays format-only and has no LLM/network code.

    ocr_lang (image inputs only): Tesseract --lang string for the OCR
    pass, e.g. "chi_sim" for Chinese slides, "deu" for German. NOT the
    translation target - this is what script/language Tesseract expects
    to SEE in the image; getting it wrong produces garbled OCR text that
    then gets "translated" into fluent-sounding nonsense (see
    _ocr_image_text's docstring). Ignored for every other file type.

    Output filename: "<stem>_<target_language><out_ext>", written next
    to the source file, or under output_dir_override if set (mirrors
    run_pipeline.py's output_dir_override convention for the AV
    pipeline). Returns the written path. Raises ValueError/RuntimeError
    on an unsupported or unreadable input - callers (GUI/CLI batch
    loops) are expected to catch and log per-file rather than let one
    bad file abort a whole batch.
    """
    src_path = Path(src_path)
    ext = src_path.suffix.lower()

    if ext in UNSUPPORTED_LEGACY_EXTENSIONS:
        modern_ext = ".docx" if ext == ".doc" else ".pptx"
        raise ValueError(
            f"{ext} (legacy binary Word/PowerPoint format, pre-2007) is not "
            f"supported. Please open {src_path.name} in Office and save it as "
            f"{modern_ext} first, then translate that file."
        )
    if ext not in TRANSLATABLE_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")
    if not src_path.is_file():
        raise FileNotFoundError(f"File not found: {src_path}")

    out_ext = output_suffix_for(src_path)
    out_dir = Path(output_dir_override) if output_dir_override else src_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    dest_path = out_dir / f"{src_path.stem}_{target_language}{out_ext}"

    if ext == ".txt":
        translate_txt(src_path, dest_path, translate_fn)
    elif ext == ".docx":
        translate_docx(src_path, dest_path, translate_fn)
    elif ext == ".pptx":
        translate_pptx(src_path, dest_path, translate_fn)
    elif ext in (".html", ".htm"):
        translate_html(src_path, dest_path, translate_fn)
    elif ext == ".pdf":
        translate_pdf_to_docx(src_path, dest_path, translate_fn)
    elif ext in IMAGE_EXTENSIONS:
        translate_image_to_docx(src_path, dest_path, translate_fn, ocr_lang=ocr_lang)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    log.info("Translated %s -> %s", src_path.name, dest_path.name)
    return dest_path
